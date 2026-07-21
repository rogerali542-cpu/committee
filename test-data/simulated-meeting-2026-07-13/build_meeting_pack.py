from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
FONT = "Microsoft YaHei"
BLUE = "1F4E78"
LIGHT = "D9EAF7"
GRAY = "666666"

def font(run, size=11, bold=False, color="000000"):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc.get_or_add_tcPr(); mar = tc.first_child_found_in("w:tcMar")
    if mar is None: mar = OxmlElement("w:tcMar"); tc.append(mar)
    for side, value in (("top",top),("start",start),("bottom",bottom),("end",end)):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), str(value)); el.set(qn("w:type"), "dxa"); mar.append(el)

def base_doc(title, subtitle):
    d=Document(); s=d.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
    normal=d.styles["Normal"]; normal.font.name=FONT; normal._element.rPr.rFonts.set(qn("w:eastAsia"),FONT); normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.1
    for name,size in (("Title",23),("Heading 1",16),("Heading 2",13)):
        st=d.styles[name]; st.font.name=FONT; st._element.rPr.rFonts.set(qn("w:eastAsia"),FONT); st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(BLUE)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4); font(p.add_run(title),23,True,BLUE)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(18); font(p.add_run(subtitle),11,False,GRAY)
    return d

def add_table(d, rows, widths=None, header=False):
    t=d.add_table(rows=0, cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.style="Table Grid"
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            if widths: cells[i].width=Inches(widths[i])
            cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cells[i])
            if header and ri==0: shade(cells[i],LIGHT)
            p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
            font(p.add_run(str(val)),10.5,header and ri==0,BLUE if header and ri==0 else "000000")
    return t

def heading(d,text,level=1): d.add_heading(text,level=level)
def para(d,text,bold_prefix=None):
    p=d.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix): font(p.add_run(bold_prefix),11,True); font(p.add_run(text[len(bold_prefix):]),11)
    else: font(p.add_run(text),11)

def notice():
    d=base_doc("关于召开星河家园业主委员会2026年第4次会议的通知","模拟测试材料｜编号：XHYW-TEST-202607-04")
    para(d,"各位委员：")
    para(d,"为测试会议小程序的通知、讨论、表决、录音补充、纪要生成及归档功能，定于2026年7月16日召开星河家园业主委员会2026年第4次会议。现将有关事项通知如下：")
    add_table(d,[["会议时间","2026年7月16日（星期四）19:30—20:30"],["会议地点","星河家园物业服务中心二楼会议室"],["会议形式","线下会议（小程序同步记录）"],["召集人","业委会主任 李明（虚构）"],["参会人员","委员7人；物业列席1人；记录员1人"]],[1.35,5.15])
    heading(d,"会议议题")
    add_table(d,[["序号","议题类型","议题名称","处理方式"],["1","通知","汛期地下车库防汛值守安排通报","听取通知并确认已知悉"],["2","讨论","儿童活动区开放时段与噪声管理优化方案","充分讨论并形成修改意见"],["3","表决","公共区域照明节能改造试点方案","记名表决：赞成/反对/弃权"]],[0.55,0.85,3.65,1.45],True)
    heading(d,"会前要求")
    for x in ["请各委员提前阅读《会议材料汇编》，重点查看节能改造预算与验收条件。","因故不能参会的委员，请按小程序流程提交请假或委托信息。","表决议题须在会议现场选择“赞成、反对或弃权”，不得以未操作代替弃权。","本通知及全部附件均为虚构测试数据，不用于真实决策。"]: para(d,"• "+x)
    para(d,"星河家园业主委员会（模拟）"); para(d,"2026年7月13日")
    d.save(OUT/"01-会议通知.docx")

def material():
    d=base_doc("星河家园业主委员会2026年第4次会议材料汇编","模拟测试材料｜含通知、讨论、表决三类议题")
    heading(d,"会议基本信息")
    add_table(d,[["会议时间","2026年7月16日 19:30—20:30"],["会议地点","物业服务中心二楼会议室"],["法定委员数","7人"],["预计出席","7人"],["测试重点","弃权票记录、票数校验、结果展示、纪要与归档"]],[1.35,5.15])
    heading(d,"议题一｜汛期地下车库防汛值守安排通报")
    para(d,"议题类型：通知。物业服务中心已完成3台排水泵试运行，并在B1层东、西坡道各配置防汛沙袋80只。7月15日至9月15日，遇橙色及以上暴雨预警时实行双人值守，每2小时巡查集水井、配电间及坡道排水沟。")
    para(d,"提请会议：委员确认知悉；如发现物资缺口，可在议题意见中补充，不进行表决。")
    heading(d,"议题二｜儿童活动区开放时段与噪声管理优化方案")
    para(d,"议题类型：讨论。近期收到关于晨间及午休时段活动噪声的模拟意见。建议试行开放时段为工作日08:30—12:00、14:30—20:30，周末及法定节假日09:00—12:00、14:30—20:30。")
    para(d,"讨论要点：是否保留暑期延时；午休时段是否完全关闭；告示牌设置位置；物业劝导与居民投诉的闭环时限。")
    para(d,"拟形成意见：试行30天，由物业每周汇总投诉数量及处理时长，下次会议复盘；本议题本次仅讨论，不表决。")
    heading(d,"议题三｜公共区域照明节能改造试点方案")
    para(d,"议题类型：表决。拟对1号楼地下车库及一层公共走廊共60盏照明灯进行LED及人体感应改造，预算上限18,000元，从公共收益列支。试点周期60天。")
    add_table(d,[["项目","测试方案"],["实施范围","1号楼地下车库40盏、一层公共走廊20盏"],["预算上限","人民币18,000元（含设备、安装、调试）"],["验收条件","照度不低于改造前；故障率≤2%；节电率目标≥25%"],["付款节点","验收合格后支付90%，质保满3个月支付10%"],["回退条件","连续两周出现照度投诉且整改无效时恢复常亮模式"]],[1.35,5.15])
    heading(d,"表决规则与模拟结果")
    para(d,"表决选项：赞成、反对、弃权。每名委员仅可选择一项；弃权属于有效参与结果，必须单独计数并显示。")
    add_table(d,[["委员（均为虚构）","选择","测试备注"],["李明","赞成","验证赞成票"],["王芳","赞成","验证赞成票"],["陈杰","赞成","验证赞成票"],["赵敏","赞成","验证赞成票"],["周强","反对","验证反对票"],["孙悦","弃权","重点验证弃权"],["吴涛","弃权","重点验证多张弃权"]],[2.2,1.1,3.2],True)
    para(d,"模拟统计：赞成4票、反对1票、弃权2票；已表决7人，未表决0人。按“全体委员过半数赞成”的测试口径，议案通过。最终结果以小程序实际提交数据为准。")
    heading(d,"补充录音关联说明")
    para(d,"补充录音对应议题三，内容包含预算上限、试点范围、弃权票说明和模拟表决结果。上传时建议标题填写“议题三补充说明录音”，关联至表决议题，用于测试转写、补充材料展示及纪要引用。")
    heading(d,"功能验收清单")
    for x in ["通知议题可查看并记录知悉状态。","讨论议题可录入多名委员意见，不强制出现表决按钮。","表决议题显示赞成、反对、弃权三个互斥选项。","弃权票计入已表决人数，并在统计、纪要和归档中独立显示。","会议通知、会议材料和补充录音均可上传、预览或播放。","生成纪要后，议题类型、票数和通过结论与原始记录一致。"]: para(d,"□ "+x)
    d.save(OUT/"02-会议材料汇编.docx")

if __name__ == "__main__": notice(); material()
