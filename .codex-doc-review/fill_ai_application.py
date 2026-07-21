import os
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

SRC = os.environ["SRC_DOCX"]
OUT = os.environ["OUT_DOCX"]


def set_cell(cell, text, size=9, align=WD_ALIGN_PARAGRAPH.LEFT):
    # Retain cell/table geometry while replacing placeholder copy.
    p0 = cell.paragraphs[0]
    ppr = deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    first = cell.paragraphs[0]
    for child in list(first._element):
        if child.tag != qn("w:pPr"):
            first._element.remove(child)
    for p in list(cell.paragraphs)[1:]:
        cell._element.remove(p._element)
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        p = cell.add_paragraph() if idx else first
        if idx == 0 and ppr is not None:
            p._p.insert(0, ppr)
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        r.font.size = Pt(size)
        r.font.name = "宋体"
        r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


d = Document(SRC)
t0 = d.tables[0]
t0_values = {
    (0, 1): "上海市静安区人民政府临汾路街道办事处",
    (4, 2): "社区治理业委会智能履职辅助场景",
    (5, 2): "依据本市住宅物业管理、业主大会和业主委员会规范化运作等有关要求，临汾路街道需进一步加强对业委会履职过程的服务和管理。当前会议材料录入、签到表整理、录音转写、会议纪要编制、事项提醒和资料归档等工作主要依靠人工完成，存在信息分散、重复录入、整理耗时、过程材料不易查询等问题，需要通过数字化和智能化手段提升工作规范性与办理效率。",
    (6, 2): "建设面向业委会履职工作的智能辅助场景，围绕业委会会议、业主接待、学习培训、提醒与待办、会议材料公示与留档等业务形成统一工作入口。重点建设业委会会议智能助手，实现图片、PDF及Word材料识别，会议录音转写，会议纪要初稿生成，议题、决议和待办事项提取，并由主任或经办人员审核确认；同时对接临汾路街道现有可信档案馆，完成经确认材料的规范留档。通过系统建设减少重复录入和人工整理时间，提升履职过程的完整性、可追溯性和管理效率。",
    (7, 2): "主要用户为临汾路街道辖区内各业委会主任及委员，覆盖全街道数十个业委会，具体人数待统计；街道相关管理人员可按权限查看工作进展、会议记录及归档情况。",
    (8, 2): "1. 业委会会议智能助手（包含材料识别、录音转写、会议纪要生成、议题及待办事项提取等能力）。",
    (9, 2): "无（现阶段不建设独立知识库）。",
}
for (row, col), text in t0_values.items():
    size = 8 if row in (5, 6) else 8.5
    align = WD_ALIGN_PARAGRAPH.CENTER if row in (0, 4) else WD_ALIGN_PARAGRAPH.LEFT
    set_cell(t0.cell(row, col), text, size, align)

t = d.tables[2]

# Header fields that can be completed without inventing personal information.
set_cell(t.cell(0, 1), "上海市静安区人民政府临汾路街道办事处", 9, WD_ALIGN_PARAGRAPH.CENTER)

values = {
    4: "业委会会议智能助手（流程自动化型、内容生成辅助型）",
    5: "申报拟采用区级统一模型能力：\n☑ Qwen2.5-VL-32B-Instruct-2lwvq2（材料识别）\n☑ DeepSeek-R1-Distill-Llama-70B-9grz3l（纪要生成、事项提取）\n说明：当前原型使用豆包视觉、语言及语音模型，实施时按区级平台可用模型适配。",
    6: "☑ 嵌入现有业务系统（临汾路街道业委会智能履职辅助系统）\n通过标准接口调用区级模型服务；业务系统继续负责身份权限、流程控制、人工确认和结果留档。",
    7: "文本类任务：平均输入约4,000—12,000 tokens，平均输出约1,000—3,000 tokens；最大输入约32,000 tokens、最大输出约8,000 tokens。长录音转写文本采用分段处理后汇总。",
    8: "预计并发用户数20人以内；常态并发模型任务1—3个，会议集中时峰值约5个。",
    9: "高：业委会会议召开前后及材料集中整理时段；\n普通：工作日8:30—17:30；\n空闲：夜间、周末及法定节假日（会议场景除外）。",
    10: "日均约10—30次模型调用，折算平均QPS小于0.01；会议材料集中处理时峰值QPS约1，非会议时段基本无调用。",
    11: "☑ 有条件共享级。\n原因：能力组件及通用提示词可复用，但实际输入可能包含业委会工作材料、会议录音、人员信息等内部数据，须经授权并按最小权限调用。",
    12: "无（现阶段不建设独立向量知识库）。\n系统仅在单次业务流程中使用经授权的会议材料、录音转写文本及业务表单数据作为上下文，不进行跨项目开放检索。",
    13: "1. 业委会会议通知、议题、附件、签到及表决结果等业务数据，由业委会成员在系统内录入或上传。\n2. 会议录音及转写文本，经参会人员知情并由主任发起采集。\n3. 业主接待、学习培训等业务记录，由相关人员在履职过程中录入。",
    14: "结构化业务数据—关系数据库；图片、PDF、Word、录音等文件—对象存储；AI生成结果及审核状态—关系数据库；归档材料—通过接口推送可信档案馆。",
    15: "中。数据由业务人员在实际履职过程中形成，来源明确；材料格式、录音质量和填写完整度存在差异，系统通过必填校验、识别结果人工确认及异常提示保障核心使用。",
    16: "具备。采用身份认证、角色权限、最小权限访问、传输加密、操作日志、文件访问控制、敏感信息提示及人工审核后入库等措施。",
    17: "不适用。现阶段不建设独立知识库；业务数据按照系统权限和数据安全要求管理，不作为知识库共享。",
    18: "纯API方式。业务系统通过标准接口调用区级统一模型、OCR及语音识别能力；流程编排、权限校验、人工确认和结果入库由业务系统实现。",
    19: "后台：Java 17（Spring Boot）；AI/OCR/语音服务：Python 3.x；前端：Vue 3 + Vite。",
    20: "OCR与文档解析、语音识别、录音转写、长文本摘要、会议纪要生成、议题及待办事项提取、文本润色、文件存储，以及可信档案馆接口。",
    21: "业委会会议智能助手 + 流程自动化型/内容生成辅助型 + 目标值60%。\n口径：在材料识别、录音转写、纪要初稿、议题和待办提取等可智能辅助环节中，由智能体完成初步处理并提交人工确认的环节占比；会议决策、表决结果确认和正式发布仍由人员负责。",
    22: "业委会会议智能助手：目标处理效率提升50%。\n预计材料录入、录音整理和会议纪要初稿制作时间由单次约4小时缩短至约2小时；最终结果须经主任或经办人员审核确认。",
}

for row, text in values.items():
    size = 8 if row in (5, 7, 13, 21, 22) else 8.5
    set_cell(t.cell(row, 2), text, size)

d.save(OUT)
print(OUT)
